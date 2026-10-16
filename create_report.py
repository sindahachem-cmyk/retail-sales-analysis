
import os
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

# Create the Word document
doc = Document()

# Set default font
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(11)

# Load the dataset
df = pd.read_csv("retail_sales_dataset.csv")
df["Date"] = pd.to_datetime(df["Date"])

# Calculate results
total_revenue = df["Total Amount"].sum()
total_transactions = len(df)
average_transaction = df["Total Amount"].mean()
category_sales = df.groupby("Product Category")["Total Amount"].sum()
gender_sales = df.groupby("Gender")["Total Amount"].sum()
monthly_sales = df.set_index("Date").resample("MS")["Total Amount"].sum()

# Title page
title = doc.add_heading("RETAIL SALES DATA ANALYSIS", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph(
    "A Data Analysis Project Using Python"
)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("\n")
p = doc.add_paragraph("Prepared by: Sinda Hachem")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph("Tools: Python, Pandas, Matplotlib")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph("Dataset: Retail Sales Dataset (1,000 transactions)")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# 1. Introduction
doc.add_heading("1. Introduction", level=1)
doc.add_paragraph(
    "This project analyzes retail transaction data to understand sales "
    "performance, customer purchasing patterns, and revenue distribution "
    "across product categories and genders. Python was used to explore "
    "the dataset, calculate key metrics, and create visualizations."
)

# 2. Objectives
doc.add_heading("2. Project Objectives", level=1)
objectives = [
    "Explore the dataset and examine data quality.",
    "Calculate total revenue and average transaction amount.",
    "Compare revenue across product categories.",
    "Compare revenue by gender.",
    "Examine monthly revenue trends."
]
for objective in objectives:
    doc.add_paragraph(objective, style="List Bullet")

# 3. Tools and methodology
doc.add_heading("3. Tools and Methodology", level=1)
doc.add_paragraph(
    "Python was used as the programming language. Pandas was used to "
    "load, inspect, group, and summarize the data. Matplotlib was used "
    "to generate charts. The analysis included checking missing values "
    "and duplicate rows, calculating summary statistics, grouping revenue "
    "by category and gender, and aggregating revenue by month."
)

# 4. Dataset overview
doc.add_heading("4. Dataset Overview", level=1)
doc.add_paragraph(
    f"The dataset contains {len(df):,} transactions and {len(df.columns)} "
    "columns. These include transaction ID, date, customer ID, gender, "
    "age, product category, quantity, price per unit, and total amount."
)
doc.add_paragraph(
    f"Missing values: {int(df.isnull().sum().sum())}. "
    f"Duplicate rows: {int(df.duplicated().sum())}."
)

# 5. Key performance indicators
doc.add_heading("5. Key Performance Indicators", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Light Shading Accent 1"
table.rows[0].cells[0].text = "Metric"
table.rows[0].cells[1].text = "Result"

metrics = [
    ("Total revenue", f"{total_revenue:,.2f}"),
    ("Number of transactions", f"{total_transactions:,}"),
    ("Average transaction amount", f"{average_transaction:,.2f}")
]
for metric, value in metrics:
    cells = table.add_row().cells
    cells[0].text = metric
    cells[1].text = value

# 6. Revenue by category
doc.add_heading("6. Revenue by Product Category", level=1)
doc.add_picture(
    "results/revenue_by_category.png",
    width=Inches(6)
)
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

for category, revenue in category_sales.sort_values(ascending=False).items():
    doc.add_paragraph(f"{category}: {revenue:,.2f}", style="List Bullet")

top_category = category_sales.idxmax()
doc.add_paragraph(
    f"{top_category} generated the highest revenue among the categories "
    "in this dataset. Clothing followed closely, while Beauty generated "
    "the lowest total revenue."
)

# 7. Revenue by gender
doc.add_heading("7. Revenue by Gender", level=1)
doc.add_picture(
    "results/revenue_by_gender.png",
    width=Inches(5.5)
)
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

for gender, revenue in gender_sales.sort_values(ascending=False).items():
    share = revenue / total_revenue * 100 if total_revenue else 0
    doc.add_paragraph(
        f"{gender}: {revenue:,.2f} ({share:.1f}% of revenue)",
        style="List Bullet"
    )

doc.add_paragraph(
    "Female customers contributed slightly more total revenue than male "
    "customers. This comparison describes revenue, not the number of "
    "customers or the average spending per customer."
)

# 8. Monthly revenue
doc.add_heading("8. Monthly Revenue Trend", level=1)
doc.add_picture(
    "results/monthly_revenue.png",
    width=Inches(6)
)
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

best_month = monthly_sales.idxmax()
doc.add_paragraph(
    f"The highest-revenue month in the available data was "
    f"{best_month.strftime('%B %Y')}, with revenue of "
    f"{monthly_sales.max():,.2f}."
)
doc.add_paragraph(
    "Monthly totals can help a retailer identify stronger and weaker "
    "periods and plan inventory and promotions. The chart should be "
    "reviewed before making business decisions."
)

# 9. Conclusions
doc.add_heading("9. Conclusion", level=1)
doc.add_paragraph(
    f"The analysis covered {total_transactions:,} transactions and "
    f"recorded total revenue of {total_revenue:,.2f}. The average "
    f"transaction amount was {average_transaction:,.2f}. "
    f"{top_category} was the leading product category by revenue, "
    "and female customers contributed slightly more revenue than male "
    "customers. The dataset contained no missing values or duplicate rows "
    "according to the checks performed."
)

# 10. Recommendations and limitations
doc.add_heading("10. Recommendations and Limitations", level=1)
recommendations = [
    "Investigate why Electronics generated the highest revenue and whether this pattern continues over time.",
    "Review monthly trends before planning promotions and inventory.",
    "Compare transaction counts and average transaction amounts by category and gender for deeper insights.",
    "Check profit margins and operating costs before treating revenue as business profitability.",
    "Confirm the dataset's currency and time coverage before presenting results to decision-makers."
]
for recommendation in recommendations:
    doc.add_paragraph(recommendation, style="List Bullet")

doc.add_paragraph(
    "Limitation: This report summarizes the available dataset. It does "
    "not establish why customers made purchases, and revenue alone does "
    "not measure profit."
)

# Save the report
output_file = "Retail_Sales_Analysis_Report.docx"
doc.save(output_file)

print("\nSUCCESS! Your Word report has been created.")
print("File:", os.path.abspath(output_file))
